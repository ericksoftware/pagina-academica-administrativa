# evaluaciones/views.py
from django.shortcuts import render, get_object_or_404, redirect
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.http import HttpResponse, JsonResponse, FileResponse  
from core.decorators import control_escolar_or_directivo_required, actas_required
from .models import Calificacion, Materia, ActaEvaluacion, Carrera, Unidad
from alumnos.models import Alumno
import pandas as pd
import os
from django.conf import settings
from django.template.loader import render_to_string
from weasyprint import HTML
from django.db import models
from django.views.decorators.clickjacking import xframe_options_exempt
from django.templatetags.static import static
from django.db.models import Q 
from django.core.paginator import Paginator
import math
from copy import copy
from decimal import Decimal
from pathlib import Path
from openpyxl import load_workbook
import os
import secrets

ACTA_TEMPLATE_PATH = Path(settings.BASE_DIR) / 'templates' / 'evaluaciones' / 'plantillas' / 'acta_evaluacion_base.xlsx'

def validar_clave_eliminacion(request, variable_env):
    clave_configurada = os.getenv(variable_env, '').strip()
    clave_ingresada = request.POST.get('delete_password', '').strip()

    if not clave_configurada:
        messages.error(request, f'No está configurada la variable {variable_env} en el archivo .env.')
        return False

    if not clave_ingresada:
        messages.error(request, 'Debes ingresar la clave de eliminación.')
        return False

    if not secrets.compare_digest(clave_ingresada, clave_configurada):
        messages.error(request, 'Clave de eliminación incorrecta.')
        return False

    return True

def construir_periodos_candidatos(ciclo_escolar, semestre):
    valor = (ciclo_escolar or '').strip()
    candidatos = []

    if valor:
        candidatos.append(valor)

    # Si el usuario escribe 2024-2025, intenta variantes comunes
    if '-' in valor:
        partes = [p.strip() for p in valor.split('-') if p.strip()]
        if len(partes) == 2 and all(p.isdigit() for p in partes):
            inicio, fin = partes
            candidatos.extend([
                inicio,
                fin,
                f'{inicio}-{semestre}',
                f'{fin}-{semestre}',
            ])
    elif valor.isdigit():
        candidatos.append(f'{valor}-{semestre}')

    # quitar duplicados conservando orden
    return list(dict.fromkeys(candidatos))

def codigo_carrera_completo(carrera):
    codigo = str(carrera.codigo or '').strip()
    plan = str(carrera.plan or '').strip()
    return f'{codigo}{plan}'

@control_escolar_or_directivo_required
def grade_list(request):
    """Lista de calificaciones"""
    calificaciones = Calificacion.objects.all().select_related('alumno', 'materia')
    
    # Filtrar por parámetros GET
    carrera = request.GET.get('carrera')
    semestre = request.GET.get('semestre')
    materia_id = request.GET.get('materia')
    
    if carrera:
        calificaciones = calificaciones.filter(alumno__carrera=carrera)
    if semestre:
        calificaciones = calificaciones.filter(alumno__semestre_actual=semestre)
    if materia_id:
        calificaciones = calificaciones.filter(materia_id=materia_id)
    
    materias = Materia.objects.all()
    
    context = {
        'calificaciones': calificaciones,
        'materias': materias,
        'carrera_filter': carrera,
        'semestre_filter': semestre,
        'materia_filter': materia_id,
        'page_title': 'Lista de Calificaciones'
    }
    return render(request, 'evaluaciones/grade_list.html', context)

@control_escolar_or_directivo_required
def import_grades(request):
    if request.method == 'POST':
        try:
            excel_file = request.FILES['excel_file']
            carrera = request.POST['carrera']
            semestre = request.POST['semestre']
            
            # Leer el archivo Excel
            df = pd.read_excel(excel_file)
            
            # Validar columnas requeridas
            required_columns = ['MATRICULA', 'MATERIA', 'CALIFICACION']
            for col in required_columns:
                if col not in df.columns:
                    messages.error(request, f'Falta la columna requerida: {col}')
                    return redirect('import_grades')
            
            # Procesar cada fila
            success_count = 0
            error_count = 0
            
            for index, row in df.iterrows():
                try:
                    matricula = str(row['MATRICULA']).strip()
                    nombre_materia = str(row['MATERIA']).strip()  # Cambiado a nombre en lugar de código
                    calificacion_val = float(row['CALIFICACION'])
                    
                    alumno = Alumno.objects.get(matricula=matricula)
                    materia = Materia.objects.get(nombre=nombre_materia)
                    
                    # Crear o actualizar calificación
                    calificacion, created = Calificacion.objects.update_or_create(
                        alumno=alumno,
                        materia=materia,
                        periodo=f"2025-{semestre}",
                        defaults={'calificacion': calificacion_val}
                    )
                    
                    success_count += 1
                    
                except Exception as e:
                    error_count += 1
                    print(f"Error en fila {index + 2}: {e}")
            
            messages.success(request, f'Importación completada: {success_count} registros exitosos, {error_count} errores')
            return redirect('grade_list')
            
        except Exception as e:
            messages.error(request, f'Error al importar el archivo: {str(e)}')
    
    # Si es GET, mostrar el formulario
    context = {
        'page_title': 'Importar Calificaciones'
    }
    return render(request, 'evaluaciones/import_grades.html', context)

@actas_required
def generate_transcript(request):
    """Generar acta de evaluación"""
    if request.method == 'POST':
        try:
            carrera = request.POST['carrera']
            semestre = request.POST['semestre']
            grupo = request.POST['grupo']
            materia_id = request.POST['materia']
            periodo = request.POST['periodo']
            fecha_emision = request.POST['fecha_emision']
            
            materia = Materia.objects.get(id=materia_id)
            
            # Crear el acta
            acta = ActaEvaluacion.objects.create(
                carrera_id=carrera,
                semestre=semestre,
                grupo=grupo,
                materia=materia,
                periodo=periodo,
                fecha_emision=fecha_emision,
                estado='generada'
            )
            
            # Obtener calificaciones para este acta
            calificaciones = Calificacion.objects.filter(
                alumno__carrera=carrera,
                alumno__semestre_actual=semestre,
                materia=materia,
                periodo=periodo
            ).select_related('alumno')
            
            # Generar PDF
            html_string = render_to_string('evaluaciones/transcript_template.html', {
                'acta': acta,
                'calificaciones': calificaciones,
                'total_alumnos': calificaciones.count(),
                'promedio_grupo': calificaciones.aggregate(avg=models.Avg('calificacion'))['calificacion__avg'] or 0
            })
            
            html = HTML(string=html_string)
            result = html.write_pdf()
            
            # Guardar el PDF
            pdf_path = f'actas/{acta.id}_{acta.materia.nombre.replace(" ", "_")}_{acta.grupo}.pdf'
            full_path = os.path.join(settings.MEDIA_ROOT, pdf_path)
            
            os.makedirs(os.path.dirname(full_path), exist_ok=True)
            with open(full_path, 'wb') as f:
                f.write(result)
            
            acta.archivo_pdf = pdf_path
            acta.save()
            
            messages.success(request, f'Acta generada exitosamente para {materia.nombre}')
            return redirect('view_transcript', transcript_id=acta.id)
            
        except Exception as e:
            messages.error(request, f'Error al generar el acta: {str(e)}')
    
    # Si es GET, mostrar el formulario
    materias = Materia.objects.all()
    
    context = {
        'materias': materias,
        'page_title': 'Generar Acta de Evaluación'
    }
    return render(request, 'evaluaciones/generate_transcript.html', context)

@actas_required
def view_transcript(request, transcript_id):
    """Ver un acta de evaluación específica"""
    acta = get_object_or_404(ActaEvaluacion, id=transcript_id)
    calificaciones = Calificacion.objects.filter(
        alumno__carrera=acta.carrera,
        alumno__semestre_actual=acta.semestre,
        materia=acta.materia,
        periodo=acta.periodo
    ).select_related('alumno')
    
    context = {
        'acta': acta,
        'calificaciones': calificaciones,
        'total_alumnos': calificaciones.count(),
        'promedio_grupo': calificaciones.aggregate(avg=models.Avg('calificacion'))['calificacion__avg'] or 0,
        'page_title': f'Acta de {acta.materia.nombre}'
    }
    return render(request, 'evaluaciones/view_transcript.html', context)

@actas_required
def download_transcript(request, transcript_id):
    """Descargar un acta en PDF"""
    acta = get_object_or_404(ActaEvaluacion, id=transcript_id)
    
    if acta.archivo_pdf:
        file_path = acta.archivo_pdf.path
        with open(file_path, 'rb') as f:
            response = HttpResponse(f.read(), content_type='application/pdf')
            response['Content-Disposition'] = f'attachment; filename="acta_{acta.materia.nombre.replace(" ", "_")}_{acta.grupo}.pdf"'
            return response
    
    messages.error(request, 'El archivo PDF no está disponible')
    return redirect('view_transcript', transcript_id=transcript_id)

@control_escolar_or_directivo_required
def lista_carreras(request):
    """Lista todas las carreras"""
    carreras_list = Carrera.objects.all().order_by('id')
    
    # Búsqueda por nombre o código
    search_query = request.GET.get('search', '')
    if search_query:
        carreras_list = carreras_list.filter(
            Q(nombre__icontains=search_query) |
            Q(codigo__icontains=search_query)
        )
    
    # Filtros
    tipo_filter = request.GET.get('tipo', '')
    estado_filter = request.GET.get('estado', '')
    
    if tipo_filter:
        carreras_list = carreras_list.filter(tipo_carrera=tipo_filter)
    if estado_filter:
        if estado_filter == 'activa':
            carreras_list = carreras_list.filter(activa=True)
        elif estado_filter == 'inactiva':
            carreras_list = carreras_list.filter(activa=False)
    
    # Paginación
    paginator = Paginator(carreras_list, 10)
    page_number = request.GET.get('page')
    carreras = paginator.get_page(page_number)
    
    context = {
        'carreras': carreras,
        'search_query': search_query,
        'tipo_filter': tipo_filter,
        'estado_filter': estado_filter,
        'page_title': 'Gestión de Carreras'
    }
    return render(request, 'evaluaciones/carreras/lista_carreras.html', context)

@control_escolar_or_directivo_required
def agregar_carrera(request):
    """Agregar una nueva carrera"""
    if request.method == 'POST':
        try:
            nombre = request.POST['nombre']
            tipo_carrera = request.POST['tipo_carrera']
            numero_semestres = request.POST['numero_semestres']
            plan = request.POST['plan']
            codigo = request.POST.get('codigo', '')
            numero_unidades = int(request.POST.get('numero_unidades', 2))
            
            # Crear la carrera
            carrera = Carrera.objects.create(
                nombre=nombre,
                tipo_carrera=tipo_carrera,
                numero_semestres=numero_semestres,
                plan=plan,
                codigo=codigo
            )
            
            # Crear unidades automáticamente basadas en el número especificado
            año_plan = str(carrera.plan)[-2:]
            
            for i in range(1, numero_unidades + 1):
                Unidad.objects.create(
                    carrera=carrera,
                    numero=i,
                    nombre=f'Unidad {i}',
                )
            
            messages.success(request, f'Carrera "{nombre}" agregada exitosamente con {numero_unidades} unidades')
            return redirect('lista_carreras')
            
        except Exception as e:
            messages.error(request, f'Error al agregar la carrera: {str(e)}')
    
    context = {
        'page_title': 'Agregar Carrera',
        'modo': 'agregar'
    }
    return render(request, 'evaluaciones/carreras/form_carrera.html', context)

@control_escolar_or_directivo_required
def editar_carrera(request, carrera_id):
    """Editar una carrera existente"""
    carrera = get_object_or_404(Carrera, id=carrera_id)

    if request.method == 'POST':
        try:
            numero_unidades_nuevo = int(request.POST.get('numero_unidades', carrera.unidades.count()))
            numero_unidades_actual = carrera.unidades.count()

            carrera.nombre = request.POST['nombre']
            carrera.tipo_carrera = request.POST['tipo_carrera']
            carrera.numero_semestres = request.POST['numero_semestres']
            carrera.plan = request.POST['plan']
            carrera.codigo = request.POST.get('codigo', '')
            carrera.activa = 'activa' in request.POST

            carrera.save()

            if numero_unidades_nuevo != numero_unidades_actual:
                if numero_unidades_nuevo > numero_unidades_actual:
                    for i in range(numero_unidades_actual + 1, numero_unidades_nuevo + 1):
                        Unidad.objects.create(
                            carrera=carrera,
                            numero=i,
                            nombre=f'Unidad {i}'
                        )
                    messages.info(request, f'Se agregaron {numero_unidades_nuevo - numero_unidades_actual} unidades nuevas')

                elif numero_unidades_nuevo < numero_unidades_actual:
                    unidades_a_eliminar = carrera.unidades.filter(numero__gt=numero_unidades_nuevo)
                    unidades_con_materia = []

                    for unidad in unidades_a_eliminar:
                        if hasattr(unidad, 'materia'):
                            unidades_con_materia.append(unidad.codigo)
                        else:
                            unidad.delete()

                    if unidades_con_materia:
                        messages.warning(
                            request,
                            f'No se pudieron eliminar {len(unidades_con_materia)} unidades porque tienen materia asignada: {", ".join(unidades_con_materia)}'
                        )

            messages.success(request, f'Carrera "{carrera.nombre}" actualizada exitosamente')
            return redirect('detalle_carrera', carrera_id=carrera.id)

        except Exception as e:
            messages.error(request, f'Error al actualizar la carrera: {str(e)}')

    context = {
        'carrera': carrera,
        'page_title': f'Editar {carrera.nombre}',
        'modo': 'editar'
    }
    return render(request, 'evaluaciones/carreras/form_carrera.html', context)

@control_escolar_or_directivo_required
def detalle_carrera(request, carrera_id):
    """Ver detalle de una carrera"""
    carrera = get_object_or_404(Carrera, id=carrera_id)
    unidades = carrera.unidades.all().select_related('materia').order_by('numero')

    materias_por_semestre = {}
    for semestre in range(1, carrera.numero_semestres + 1):
        materias_por_semestre[semestre] = Materia.objects.filter(
            carrera=carrera,
            semestre=semestre
        ).select_related('unidad')

    context = {
        'carrera': carrera,
        'unidades': unidades,
        'materias_por_semestre': materias_por_semestre,
        'page_title': f'Detalle de {carrera.nombre}'
    }
    return render(request, 'evaluaciones/carreras/detalle_carrera.html', context)

@control_escolar_or_directivo_required
def eliminar_carrera(request, carrera_id):
    carrera = get_object_or_404(Carrera, id=carrera_id)

    if request.method == 'POST':
        if not validar_clave_eliminacion(request, 'del_ca_pass'):
            return redirect('detalle_carrera', carrera_id=carrera.id)

        try:
            nombre = carrera.nombre
            carrera.delete()
            messages.success(request, f'Carrera {nombre} eliminada exitosamente')
            return redirect('lista_carreras')
        except Exception as e:
            messages.error(request, f'Error al eliminar la carrera: {str(e)}')
            return redirect('detalle_carrera', carrera_id=carrera.id)

    return redirect('detalle_carrera', carrera_id=carrera.id)

@control_escolar_or_directivo_required
def agregar_materia(request, carrera_id):
    """Agregar una materia a una carrera"""
    carrera = get_object_or_404(Carrera, id=carrera_id)

    if request.method == 'POST':
        try:
            nombre = request.POST['nombre'].strip()
            semestre = int(request.POST['semestre'])
            unidad_id = request.POST.get('unidad')
            parciales = int(request.POST.get('parciales', 3))
            creditos = float(request.POST.get('creditos', 8.0))
            horas = int(request.POST.get('horas', 64))

            if not unidad_id:
                messages.error(request, 'Debes seleccionar una unidad.')
                return redirect('detalle_carrera', carrera_id=carrera_id)

            unidad = get_object_or_404(Unidad, id=unidad_id, carrera=carrera)

            if hasattr(unidad, 'materia'):
                messages.error(request, f'La unidad {unidad.codigo} ya está asignada a la materia "{unidad.materia.nombre}".')
                return redirect('detalle_carrera', carrera_id=carrera_id)

            Materia.objects.create(
                carrera=carrera,
                unidad=unidad,
                nombre=nombre,
                semestre=semestre,
                parciales=parciales,
                creditos=creditos,
                horas=horas
            )

            messages.success(request, f'Materia "{nombre}" agregada exitosamente')

        except Exception as e:
            messages.error(request, f'Error al agregar la materia: {str(e)}')

    return redirect('detalle_carrera', carrera_id=carrera_id)

@control_escolar_or_directivo_required
def eliminar_materia(request, materia_id):
    """Eliminar una materia"""
    materia = get_object_or_404(Materia, id=materia_id)
    carrera_id = materia.carrera.id
    
    if request.method == 'POST':
        try:
            nombre = materia.nombre
            materia.delete()
            messages.success(request, f'Materia "{nombre}" eliminada exitosamente')
        except Exception as e:
            messages.error(request, f'Error al eliminar la materia: {str(e)}')
    
    return redirect('detalle_carrera', carrera_id=carrera_id)

@control_escolar_or_directivo_required
def editar_materia(request, materia_id):
    """Editar una materia existente"""
    materia = get_object_or_404(Materia, id=materia_id)

    if request.method == 'POST':
        try:
            nueva_unidad_id = request.POST.get('unidad')
            if not nueva_unidad_id:
                messages.error(request, 'Debes seleccionar una unidad.')
                return redirect('detalle_carrera', carrera_id=materia.carrera.id)

            nueva_unidad = get_object_or_404(Unidad, id=nueva_unidad_id, carrera=materia.carrera)

            materia_existente = Materia.objects.filter(unidad=nueva_unidad).exclude(pk=materia.pk).first()
            if materia_existente:
                messages.error(
                    request,
                    f'La unidad {nueva_unidad.codigo} ya está asignada a la materia "{materia_existente.nombre}".'
                )
                return redirect('detalle_carrera', carrera_id=materia.carrera.id)

            materia.nombre = request.POST['nombre'].strip()
            materia.semestre = int(request.POST['semestre'])
            materia.unidad = nueva_unidad
            materia.parciales = int(request.POST.get('parciales', 3))
            materia.creditos = float(request.POST.get('creditos', 8.0))
            materia.horas = int(request.POST.get('horas', 64))
            materia.activa = 'activa' in request.POST
            materia.save()

            messages.success(request, f'Materia "{materia.nombre}" actualizada exitosamente')

        except Exception as e:
            messages.error(request, f'Error al actualizar la materia: {str(e)}')

    return redirect('detalle_carrera', carrera_id=materia.carrera.id)

@control_escolar_or_directivo_required
def cambiar_estado_materia(request, materia_id):
    """Cambiar el estado activo/inactivo de una materia"""
    materia = get_object_or_404(Materia, id=materia_id)
    
    if request.method == 'POST':
        try:
            materia.activa = not materia.activa
            materia.save()
            
            estado = "activada" if materia.activa else "desactivada"
            messages.success(request, f'Materia "{materia.nombre}" {estado} exitosamente')
        except Exception as e:
            messages.error(request, f'Error al cambiar el estado de la materia: {str(e)}')
    
    return redirect('detalle_carrera', carrera_id=materia.carrera.id)

@control_escolar_or_directivo_required
def obtener_datos_materia(request, materia_id):
    """Obtener datos de una materia para edición (AJAX)"""
    try:
        materia = get_object_or_404(Materia, id=materia_id)

        return JsonResponse({
            'nombre': materia.nombre,
            'semestre': materia.semestre,
            'unidad_id': materia.unidad_id,
            'parciales': materia.parciales,
            'creditos': float(materia.creditos),
            'horas': materia.horas,
            'activa': materia.activa
        })
    except Exception as e:
        return JsonResponse({'error': str(e)}, status=500)
    
# evaluaciones/views.py (agregar temporalmente para debugging)
@control_escolar_or_directivo_required
def debug_materia(request, materia_id):
    """Vista temporal para debugging"""
    materia = get_object_or_404(Materia, id=materia_id)
    
    print("=== DEBUG MATERIA ===")
    print(f"Materia: {materia.nombre}")
    print(f"Unidades asignadas: {list(materia.unidades.values_list('id', 'codigo'))}")
    print("=====================")
    
    return JsonResponse({
        'debug': 'ok',
        'materia': materia.nombre,
        'unidades': list(materia.unidades.values_list('id', flat=True))
    })

@control_escolar_or_directivo_required
def asignar_materias_unidades_sin_asignar(request, carrera_id):
    messages.info(
        request,
        'Esta acción ya no es necesaria. Ahora cada materia se crea con una sola unidad obligatoria.'
    )
    return redirect('detalle_carrera', carrera_id=carrera_id)


def normalizar_turno(turno):
    valor = (turno or '').strip().lower()

    if valor in {'m', 'matutino'}:
        return 'MATUTINO', 'M'

    if valor in {'v', 'vespertino'}:
        return 'VESPERTINO', 'V'

    raise ValueError('El turno debe ser Matutino o Vespertino.')


def nombre_completo_alumno(alumno):
    metodo = getattr(alumno, 'nombre_completo', None)

    if callable(metodo):
        try:
            return str(metodo()).strip().upper()
        except TypeError:
            pass

    partes = [
        getattr(alumno, 'apellido_paterno', '') or '',
        getattr(alumno, 'apellido_materno', '') or '',
        getattr(alumno, 'nombre', '') or '',
    ]
    return ' '.join([p.strip() for p in partes if p.strip()]).upper()


def valor_numerico_o_none(valor):
    if valor is None:
        return None

    valor = Decimal(str(valor))
    if valor == valor.to_integral_value():
        return int(valor)

    return float(valor)


def obtener_alumnos_para_acta(carrera, semestre, turno_texto, turno_letra):
    filtros_turno = (
        Q(turno__iexact=turno_texto) |
        Q(turno__iexact=turno_texto.lower()) |
        Q(turno__iexact=turno_letra) |
        Q(turno__iexact=turno_letra.lower())
    )

    return Alumno.objects.filter(
        carrera=carrera,
        semestre_actual=semestre
    ).filter(
        filtros_turno
    ).order_by(
        'apellido_paterno',
        'apellido_materno',
        'nombre'
    )


def construir_filas_acta(carrera, semestre, turno, materia, ciclo_escolar):
    turno_texto, turno_letra = normalizar_turno(turno)

    alumnos = list(
        obtener_alumnos_para_acta(carrera, semestre, turno_texto, turno_letra)
    )

    if not alumnos:
        return [], turno_texto, turno_letra

    periodos_candidatos = construir_periodos_candidatos(ciclo_escolar, semestre)

    calificaciones = Calificacion.objects.filter(
        alumno__in=alumnos,
        unidad=materia.unidad,
        periodo__in=periodos_candidatos
    )

    mapa = {}
    for calificacion in calificaciones:
        llave = (
            calificacion.alumno_id,
            calificacion.tipo_calificacion,
            calificacion.numero_parcial
        )
        mapa[llave] = valor_numerico_o_none(calificacion.calificacion)

    filas = []
    for alumno in alumnos:
        unidad_1 = mapa.get((alumno.id, 'parcial', 1))
        unidad_2 = mapa.get((alumno.id, 'parcial', 2))
        evidencia_final = mapa.get((alumno.id, 'evidencia_final', None))

        filas.append({
            'alumno_id': alumno.id,
            'codigo_alumno': alumno.id,
            'matricula': alumno.matricula,
            'nombre_completo': nombre_completo_alumno(alumno),
            'unidad_1': unidad_1,
            'unidad_2': unidad_2,
            'evidencia_final': evidencia_final,
        })

    return filas, turno_texto, turno_letra


def porcentaje_asistencia_excel(inasistencias, total_horas):
    if not total_horas:
        return 0
    return math.trunc(((total_horas - inasistencias) * 100) / total_horas)


def evaluacion_global_excel(unidad_1, unidad_2, evidencia_final, porcentaje_asistencia):
    if porcentaje_asistencia < 85:
        return None

    parciales = [v for v in [unidad_1, unidad_2] if v is not None]
    promedio_parciales = (sum(parciales) / len(parciales)) if parciales else 0
    evidencia = evidencia_final if evidencia_final is not None else 0

    return math.trunc((promedio_parciales / 2) + (evidencia / 2))


def formula_evaluacion_global(row_num):
    return (
        f'=IF(K{row_num}<85,"No%",'
        f'TRUNC((IF(ISNUMBER(E{row_num}),E{row_num},0)+IF(ISNUMBER(F{row_num}),F{row_num},0))'
        f'/IF(COUNT(D{row_num}:F{row_num})=0,1,COUNT(D{row_num}:F{row_num}))/2+'
        f'(IF(ISNUMBER(G{row_num}),G{row_num},0)/2),0))'
    )


def formula_porcentaje_inasistencia(row_num):
    return f'=TRUNC((((J$6-IF(ISNUMBER(J{row_num}),J{row_num},0))*100)/J$6),0)'


def copiar_estilo_fila(ws, fila_origen, fila_destino, max_col=11):
    for col in range(1, max_col + 1):
        origen = ws.cell(fila_origen, col)
        destino = ws.cell(fila_destino, col)

        if origen.has_style:
            destino._style = copy(origen._style)
        if origen.number_format:
            destino.number_format = origen.number_format
        if origen.font:
            destino.font = copy(origen.font)
        if origen.fill:
            destino.fill = copy(origen.fill)
        if origen.border:
            destino.border = copy(origen.border)
        if origen.alignment:
            destino.alignment = copy(origen.alignment)
        if origen.protection:
            destino.protection = copy(origen.protection)

    ws.row_dimensions[fila_destino].height = ws.row_dimensions[fila_origen].height


def generar_archivo_excel_acta(acta, filas):
    if not ACTA_TEMPLATE_PATH.exists():
        raise FileNotFoundError(
            'No se encontró la plantilla Excel del acta en templates/evaluaciones/plantillas/acta_evaluacion_base.xlsx'
        )

    wb = load_workbook(ACTA_TEMPLATE_PATH)
    ws = wb.active

    # Forzar recálculo al abrir en Excel
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True

    # Encabezado
    ws['D2'] = acta.periodo
    ws['C3'] = codigo_carrera_completo(acta.carrera)
    ws['D3'] = acta.carrera.nombre
    ws['C4'] = acta.materia.unidad.codigo
    ws['D4'] = acta.materia.nombre
    ws['C5'] = acta.docente_identificador
    ws['D5'] = acta.docente_nombre
    ws['C6'] = acta.semestre
    ws['D6'] = acta.turno
    ws['E6'] = acta.grupo
    ws['J6'] = acta.total_horas

    fila_inicio = 8
    fila_fin_base = 37
    fila_min_base = 38
    filas_base = fila_fin_base - fila_inicio + 1

    extra = max(0, len(filas) - filas_base)

    if extra > 0:
        ws.insert_rows(fila_min_base, amount=extra)
        for fila in range(fila_min_base, fila_min_base + extra):
            copiar_estilo_fila(ws, fila_fin_base, fila, max_col=11)

    fila_fin_datos = fila_fin_base + extra
    fila_min = fila_min_base + extra
    fila_max = fila_min + 1

    # Limpiar y reconstruir área de alumnos
    for fila in range(fila_inicio, fila_fin_datos + 1):
        ws[f'A{fila}'] = fila - fila_inicio + 1
        ws[f'B{fila}'] = None
        ws[f'C{fila}'] = None
        ws[f'D{fila}'] = None
        ws[f'E{fila}'] = None
        ws[f'F{fila}'] = None
        ws[f'G{fila}'] = None
        ws[f'H{fila}'] = formula_evaluacion_global(fila)
        ws[f'J{fila}'] = None
        ws[f'K{fila}'] = formula_porcentaje_inasistencia(fila)

    # Llenar filas reales
    for index, fila in enumerate(filas, start=fila_inicio):
        ws[f'B{index}'] = fila['codigo_alumno']
        ws[f'C{index}'] = fila['matricula']
        ws[f'D{index}'] = fila['nombre_completo']
        ws[f'E{index}'] = fila['unidad_1']
        ws[f'F{index}'] = fila['unidad_2']
        ws[f'G{index}'] = fila['evidencia_final']
        ws[f'J{index}'] = fila['inasistencias']

    # Pie
    ws[f'D{fila_min}'] = 'Valor mínimo aceptado'
    ws[f'E{fila_min}'] = 5
    ws[f'F{fila_min}'] = 5
    ws[f'G{fila_min}'] = 5
    ws[f'J{fila_min}'] = 0

    ws[f'D{fila_max}'] = 'Valor máximo aceptado'
    ws[f'E{fila_max}'] = 10
    ws[f'F{fila_max}'] = 10
    ws[f'G{fila_max}'] = 10
    ws[f'J{fila_max}'] = acta.total_horas

    nombre_archivo = (
        f'Califica_{acta.periodo}_{acta.carrera.codigo}_'
        f'{acta.semestre}{acta.grupo}_{acta.turno}_{acta.materia.unidad.codigo}.xlsx'
    ).replace(' ', '_')

    ruta_relativa = f'actas/{nombre_archivo}'
    ruta_absoluta = os.path.join(settings.MEDIA_ROOT, ruta_relativa)

    os.makedirs(os.path.dirname(ruta_absoluta), exist_ok=True)
    wb.save(ruta_absoluta)

    return ruta_relativa

@actas_required
def acta_list(request):
    """Lista de todas las actas de evaluación"""
    actas_list = ActaEvaluacion.objects.all().select_related('carrera', 'materia').order_by('-fecha_generacion')
    
    # Búsqueda
    search_query = request.GET.get('search', '')
    if search_query:
        actas_list = actas_list.filter(
            Q(materia__nombre__icontains=search_query) |
            Q(carrera__nombre__icontains=search_query) |
            Q(grupo__icontains=search_query) |
            Q(periodo__icontains=search_query)
        )
    
    # Filtros
    carrera_filter = request.GET.get('carrera', '')
    semestre_filter = request.GET.get('semestre', '')
    estado_filter = request.GET.get('estado', '')
    
    if carrera_filter:
        actas_list = actas_list.filter(carrera_id=carrera_filter)
    if semestre_filter:
        actas_list = actas_list.filter(semestre=semestre_filter)
    if estado_filter:
        actas_list = actas_list.filter(estado=estado_filter)
    
    carreras = Carrera.objects.all()
    carrera_selected = None
    if carrera_filter:
        carrera_selected = Carrera.objects.filter(id=carrera_filter).first()
    
    # Paginación
    paginator = Paginator(actas_list, 10)
    page_number = request.GET.get('page')
    actas = paginator.get_page(page_number)
    
    context = {
        'actas': actas,
        'carreras': carreras,
        'carrera_selected': carrera_selected,
        'search_query': search_query,
        'carrera_filter': carrera_filter,
        'semestre_filter': semestre_filter,
        'estado_filter': estado_filter,
        'page_title': 'Lista de Actas de Evaluación'
    }
    return render(request, 'evaluaciones/acta_list.html', context)

@actas_required
def acta_preview_data(request):
    """Devuelve datos previos del acta para llenar el formulario dinámicamente."""
    try:
        carrera_id = request.GET.get('carrera')
        semestre = request.GET.get('semestre')
        materia_id = request.GET.get('materia')
        turno = request.GET.get('turno')
        ciclo_escolar = request.GET.get('ciclo_escolar')

        if not all([carrera_id, semestre, materia_id, turno, ciclo_escolar]):
            return JsonResponse(
                {'error': 'Debes indicar carrera, semestre, materia, turno y ciclo escolar.'},
                status=400
            )

        carrera = get_object_or_404(Carrera, id=carrera_id)
        materia = get_object_or_404(
            Materia.objects.select_related('unidad', 'carrera'),
            id=materia_id,
            carrera=carrera
        )

        semestre_int = int(semestre)
        if materia.semestre != semestre_int:
            return JsonResponse(
                {'error': 'La materia seleccionada no pertenece al semestre indicado.'},
                status=400
            )

        filas, turno_texto, turno_letra = construir_filas_acta(
            carrera=carrera,
            semestre=semestre_int,
            turno=turno,
            materia=materia,
            ciclo_escolar=ciclo_escolar
        )

        return JsonResponse({
            'carrera': {
                'id': carrera.id,
                'codigo': f'{carrera.codigo}{carrera.plan}',
                'nombre': carrera.nombre,
                'plan': carrera.plan,
            },
            'materia': {
                'id': materia.id,
                'codigo': materia.unidad.codigo,
                'nombre': materia.nombre,
                'horas': materia.horas,
                'semestre': materia.semestre,
            },
            'turno': {
                'texto': turno_texto,
                'letra': turno_letra,
            },
            'alumnos': filas,
            'total_alumnos': len(filas),
        })
    except ValueError as e:
        return JsonResponse({'error': str(e)}, status=400)
    except Exception as e:
        return JsonResponse({'error': str(e)}, status=500)

@actas_required
def generate_acta(request):
    """Generar una nueva acta de evaluación en Excel."""
    carreras = Carrera.objects.filter(activa=True).order_by('nombre')

    if request.method == 'POST':
        try:
            carrera_id = request.POST.get('carrera')
            semestre = int(request.POST.get('semestre'))
            materia_id = request.POST.get('materia')
            turno = request.POST.get('turno')
            ciclo_escolar = request.POST.get('ciclo_escolar', '').strip()
            docente_identificador = request.POST.get('docente_identificador', '').strip()
            docente_nombre = request.POST.get('docente_nombre', '').strip().upper()
            fecha_emision = request.POST.get('fecha_emision')

            if not all([
                carrera_id,
                semestre,
                materia_id,
                turno,
                ciclo_escolar,
                docente_identificador,
                docente_nombre,
                fecha_emision,
            ]):
                raise ValueError('Todos los campos obligatorios deben estar completos.')

            carrera = get_object_or_404(Carrera, id=carrera_id)
            materia = get_object_or_404(
                Materia.objects.select_related('unidad', 'carrera'),
                id=materia_id,
                carrera=carrera
            )

            if materia.semestre != semestre:
                raise ValueError('La materia seleccionada no pertenece al semestre elegido.')

            filas, turno_texto, turno_letra = construir_filas_acta(
                carrera=carrera,
                semestre=semestre,
                turno=turno,
                materia=materia,
                ciclo_escolar=ciclo_escolar
            )

            if not filas:
                raise ValueError('No se encontraron alumnos para la carrera, semestre y turno seleccionados.')

            for fila in filas:
                campo = f'inasistencias_{fila["alumno_id"]}'
                valor = (request.POST.get(campo, '0') or '0').strip()

                try:
                    inasistencias = int(valor)
                except ValueError:
                    raise ValueError(f'Las inasistencias del alumno {fila["nombre_completo"]} deben ser numéricas.')

                if inasistencias < 0 or inasistencias > materia.horas:
                    raise ValueError(
                        f'Las inasistencias del alumno {fila["nombre_completo"]} deben estar entre 0 y {materia.horas}.'
                    )

                fila['inasistencias'] = inasistencias

            globales_numericas = []
            alumnos_evaluados = 0

            for fila in filas:
                porcentaje = porcentaje_asistencia_excel(fila['inasistencias'], materia.horas)
                global_excel = evaluacion_global_excel(
                    fila['unidad_1'],
                    fila['unidad_2'],
                    fila['evidencia_final'],
                    porcentaje
                )

                if any(v is not None for v in [fila['unidad_1'], fila['unidad_2'], fila['evidencia_final']]):
                    alumnos_evaluados += 1

                if global_excel is not None:
                    globales_numericas.append(global_excel)

            promedio_grupo = (
                sum(globales_numericas) / len(globales_numericas)
                if globales_numericas else 0
            )

            acta, creada = ActaEvaluacion.objects.update_or_create(
                materia=materia,
                grupo=turno_letra,
                periodo=ciclo_escolar,
                defaults={
                    'carrera': carrera,
                    'semestre': semestre,
                    'turno': turno_texto,
                    'fecha_emision': fecha_emision,
                    'estado': 'generada',
                    'generada_por': request.user,
                    'docente_identificador': docente_identificador,
                    'docente_nombre': docente_nombre,
                    'total_horas': materia.horas,
                    'total_alumnos': len(filas),
                    'alumnos_evaluados': alumnos_evaluados,
                    'promedio_grupo': promedio_grupo,
                }
            )

            # AQUÍ VA
            if acta.archivo_excel and acta.archivo_excel.name:
                acta.archivo_excel.delete(save=False)

            ruta_excel = generar_archivo_excel_acta(acta, filas)
            acta.archivo_excel.name = ruta_excel
            acta.save()

            messages.success(request, f'Acta de evaluación generada exitosamente para {materia.nombre}')
            return redirect('view_acta', acta_id=acta.id)

        except Exception as e:
            messages.error(request, f'Error al generar el acta: {str(e)}')

    context = {
        'carreras': carreras,
        'page_title': 'Generar Acta de Evaluación'
    }
    return render(request, 'evaluaciones/generate_acta.html', context)

@actas_required
def view_acta(request, acta_id):
    """Ver una acta específica."""
    acta = get_object_or_404(
        ActaEvaluacion.objects.select_related('carrera', 'materia', 'materia__unidad'),
        id=acta_id
    )

    context = {
        'acta': acta,
        'page_title': f'Acta de {acta.materia.nombre}'
    }
    return render(request, 'evaluaciones/view_acta.html', context)

@actas_required
def download_acta(request, acta_id):
    """Descargar una acta en Excel."""
    acta = get_object_or_404(ActaEvaluacion, id=acta_id)

    if acta.archivo_excel and acta.archivo_excel.name:
        try:
            return FileResponse(
                acta.archivo_excel.open('rb'),
                as_attachment=True,
                filename=os.path.basename(acta.archivo_excel.name),
                content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
            )
        except Exception as e:
            messages.error(request, f'Error al descargar el archivo: {str(e)}')
            return redirect('view_acta', acta_id=acta.id)

    messages.error(request, 'El archivo Excel no está disponible')
    return redirect('view_acta', acta_id=acta.id)

@xframe_options_exempt
@actas_required
def view_acta_pdf(request, acta_id):
    """Vista especial para mostrar PDF en iframe - Similar a view_pdf"""
    acta = get_object_or_404(ActaEvaluacion, id=acta_id)
    
    if acta.archivo_pdf and acta.archivo_pdf.name:
        try:
            response = FileResponse(
                acta.archivo_pdf.open(),
                content_type='application/pdf'
            )
            response['Content-Disposition'] = f'inline; filename="acta_{acta.materia.nombre.replace(" ", "_")}_{acta.grupo}.pdf"'
            return response
        except Exception as e:
            return HttpResponse("Error al cargar el PDF", status=500)
    else:
        return HttpResponse("PDF no disponible", status=404)
    
@actas_required
def obtener_materias_por_carrera(request, carrera_id):
    """Obtener materias activas por carrera y semestre (AJAX)."""
    try:
        carrera = get_object_or_404(Carrera, id=carrera_id)
        semestre = request.GET.get('semestre')

        materias = Materia.objects.filter(
            carrera=carrera,
            activa=True
        ).select_related('unidad').order_by('semestre', 'nombre')

        if semestre:
            materias = materias.filter(semestre=semestre)

        materias_data = []
        for materia in materias:
            materias_data.append({
                'id': materia.id,
                'nombre': materia.nombre,
                'semestre': materia.semestre,
                'codigo': materia.unidad.codigo if materia.unidad_id else 'N/A',
                'horas': materia.horas,
            })

        return JsonResponse({'materias': materias_data})
    except Exception as e:
        return JsonResponse({'error': str(e)}, status=500)
